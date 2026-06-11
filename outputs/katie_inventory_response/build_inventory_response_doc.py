from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/yakub/Desktop/iblotmark/outputs/katie_inventory_response/Inventory Bin Count Automation - Recommendation and Pilot Plan.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(17, 24, 39)
MUTED = RGBColor(75, 85, 99)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.15)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def format_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                set_paragraph_spacing(p, after=0, line=1.10)
                for run in p.runs:
                    set_run_font(run, size=10, color=INK)
            if header and row_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, size=10, color=INK, bold=True)


def add_key_value_rows(doc, rows):
    for label, value in rows:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2, line=1.10)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=11, color=INK, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=11, color=INK)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(12)
    subtitle.font.italic = True
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(12)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Inventory Bin Count Automation | Pilot Recommendation")
    set_run_font(r, size=9, color=MUTED)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("Inventory Bin Count Automation")
    set_run_font(title_run, size=22, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("Plain-English recommendation and next-quarter pilot plan")
    set_run_font(subtitle_run, size=12, color=MUTED, italic=True)
    add_key_value_rows(
        doc,
        [
            ("Prepared for", "Katie Hobbs"),
            ("Prepared by", "Jacob / iBolt team"),
            ("Date", "May 14, 2026"),
            ("Purpose", "Respond to the request for a faster, lower-error inventory bin counting process using barcodes, weight, and simple automation."),
        ],
    )

    add_callout(
        doc,
        "Recommendation",
        "Yes, this is worth piloting. Start with a simple barcode + spreadsheet workflow first, then add direct scale integration only after the team proves the process works on real bins.",
    )

    add_heading(doc, "What This Would Do", 1)
    add_para(
        doc,
        "The system would remove most manual lookup and calculator work from inventory counts. Each bin would have a barcode. When someone scans the bin, the system identifies the SKU, product name, unit weight, and empty bin weight. The user then enters the gross weight from the scale, and the sheet or app calculates the bin quantity.",
    )
    add_para(
        doc,
        "The counting formula stays exactly the same as the current process, but it is handled automatically:",
        after=2,
    )
    add_callout(
        doc,
        "Formula",
        "(Total bin weight in oz - empty bin weight in oz) / individual part weight in oz = calculated quantity.",
    )

    add_heading(doc, "Recommended Workflow", 1)
    steps = [
        ("Label the bin", "Print a barcode label for each bin. The label should show the bin ID, SKU, product name, and a scannable Code 128 barcode."),
        ("Scan the bin", "The scanner enters the bin ID automatically, like a keyboard. No typing the SKU by hand."),
        ("Weigh the bin", "Put the bin on the scale and enter the total weight in pounds."),
        ("Review the result", "The system calculates quantity, rounds to the nearest whole unit, and flags rows that need a reweigh or manual check."),
        ("Export counts", "After review, export the counted quantities to the inventory system or final count workbook."),
    ]
    for i, (label, text) in enumerate(steps, start=1):
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=6, line=1.167)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=11, color=INK, bold=True)
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK)

    add_heading(doc, "What To Buy", 1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    for idx, value in enumerate(["Item", "Buy Now?", "Budget", "Reason"]):
        hdr[idx].text = value
    rows = [
        ("USB 2D barcode scanner", "Yes", "$160-$250 each", "Reliable scanning; works as keyboard input with spreadsheets or a simple app."),
        ("Label printer", "Yes", "$250-$1,100", "Needed to print readable bin labels with barcode plus human-readable text."),
        ("Thermal labels", "Yes", "$20-$80 per roll", "Use 2x1, 3x1, or 4x2 labels depending on bin size and viewing distance."),
        ("New scale", "Not first", "$0 initially", "Use the existing scale for the pilot. Buy an integrated USB/RS-232 scale only if the pilot works."),
        ("Scale integration", "Phase 2", "$450-$700+ scale, plus setup", "Direct scale capture is useful, but not required to prove the workflow."),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value
    format_table(table, [2300, 1300, 1600, 4160])

    add_heading(doc, "Important Accuracy Notes", 1)
    add_para(
        doc,
        "This should be treated as a counting aid, not a magic count. It will be very useful when each bin contains one SKU and the unit weight is trustworthy. The biggest accuracy risks are small hardware, mixed-SKU bins, inconsistent packaging, and scale resolution.",
    )
    notes = [
        "For small parts, verify unit weight by weighing 50 or 100 pieces, then dividing by the sample count.",
        "Keep the default empty bin weight at 56 oz, but record actual bin tare if some bins are meaningfully different.",
        "Rows with unusual fractional results should be flagged for reweigh or manual count.",
        "Do not rely on AI to do the math. AI can help with setup and review, but the count should be deterministic and formula-based.",
    ]
    for item in notes:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=6, line=1.167)
        r = p.add_run(item)
        set_run_font(r, size=11, color=INK)

    add_heading(doc, "Pilot Plan", 1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    for idx, value in enumerate(["Timing", "Task", "Outcome"]):
        hdr[idx].text = value
    rows = [
        ("Week 1", "Confirm SKU list, unit weights, bin IDs, and current scale readability.", "Clean source data before labels are printed."),
        ("Week 2", "Print labels and load the spreadsheet or simple app with SKU/bin data.", "Ready-to-test counting station."),
        ("Week 3", "Pilot 25-50 bins across light, medium, and heavy parts.", "Measure speed, accuracy, and operator friction."),
        ("Week 4", "Fix issues and decide whether to keep spreadsheet MVP or build a small app.", "Clear go/no-go decision before broader rollout."),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value
    format_table(table, [1500, 4300, 3560])

    add_heading(doc, "Decision", 1)
    add_para(
        doc,
        "The best next step is a low-cost pilot, not a large software purchase. A barcode scanner, labels, and the prototype workbook are enough to validate whether this saves time and reduces errors. If the pilot is successful, the next version can be a small web app with direct scale capture, audit logs, and export back to the inventory system.",
    )

    add_heading(doc, "Files Included In The Packet", 1)
    included = [
        ("Inventory Bin Count Automation - Recommendation and Pilot Plan.docx", "This plain-English memo."),
        ("inventory_bin_count_template.xlsx", "Spreadsheet prototype that demonstrates the count logic."),
        ("Reply Email Draft.md", "Copy/paste response email to Katie."),
    ]
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "File"
    hdr[1].text = "Purpose"
    for row_data in included:
        row = table.add_row().cells
        row[0].text = row_data[0]
        row[1].text = row_data[1]
    format_table(table, [4700, 4660])

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
